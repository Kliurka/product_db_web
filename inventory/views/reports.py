from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.template.loader import render_to_string
from django.utils import timezone

from weasyprint import HTML

from inventory.models import Order
from inventory.utils.product_filters import get_filtered_products
from inventory.utils.permissions import role_required


@login_required
@role_required("admin", "manager")
def order_pdf(request, order_code):

    order = get_object_or_404(
        Order.objects.select_related(
            "customer",
            "created_by",
            "updated_by",
        ),
        order_code=order_code,
    )

    items = (
        order.items
        .select_related("product")
        .all()
    )

    company = {
        "name": "UAB Akmens gaminiai",
        "company_code": "123456789",
        "vat_code": "LT123456789",
        "address": "Gatvė 1, Panevėžys, Lietuva",
        "phone": "+370 600 00000",
        "email": "info@example.lt",
        "website": "www.example.lt",
        "bank": "Swedbank",
        "iban": "LT00 0000 0000 0000 0000",
    }

    html_string = render_to_string(
        "inventory/pdf/order_pdf.html",
        {
            "company": company,
            "order": order,
            "items": items,
            "generated_at": timezone.localtime(),
        },
    )

    pdf_file = HTML(
        string=html_string,
        base_url=request.build_absolute_uri("/"),
    ).write_pdf()

    response = HttpResponse(
        pdf_file,
        content_type="application/pdf",
    )

    response["Content-Disposition"] = (
        f'inline; filename="order-{order.order_code}.pdf"'
    )

    return response


@login_required
@role_required("admin", "manager", "worker")
def production_pdf(request, order_code):

    order = get_object_or_404(
        Order.objects.select_related(
            "customer",
        ),
        order_code=order_code,
    )

    items = (
        order.items
        .select_related("product")
        .all()
    )

    company = {
        "name": "STONE FACTORY",
    }

    html_string = render_to_string(
        "inventory/pdf/production_pdf.html",
        {
            "company": company,
            "order": order,
            "items": items,
        },
    )

    pdf = HTML(
        string=html_string,
        base_url=request.build_absolute_uri("/"),
    ).write_pdf()

    response = HttpResponse(
        pdf,
        content_type="application/pdf",
    )

    response["Content-Disposition"] = (
        f'inline; filename="production-{order.order_code}.pdf"'
    )

    return response


@login_required
@role_required("admin", "manager", "worker", "viewer")
def product_list_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    products = get_filtered_products(request).order_by("code")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"

    headers = [
        "Code",
        "Name",
        "Type",
        "Texture",
        "Dim X",
        "Dim Y",
        "Description",
        "Storage Location",
        "Price",
        "Status",
        "Created At",
        "Created By",
        "Updated At",
        "Updated By",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for product in products:
        worksheet.append([
            product.code,
            product.name,
            str(product.type or ""),
            str(product.texture or ""),
            product.dimension_x,
            product.dimension_y,
            product.description or "",
            str(product.storage_location or ""),
            product.price,
            product.get_status_display(),
            product.created_at.replace(tzinfo=None) if product.created_at else None,
            str(product.created_by or ""),
            product.updated_at.replace(tzinfo=None) if product.updated_at else None,
            str(product.updated_by or ""),
        ])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    widths = [
        16, 26, 18, 18, 12, 12, 34,
        18, 12, 14, 18, 16, 18, 16,
    ]

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="filtered-products.xlsx"'
    )

    return response


@login_required
@role_required("admin", "manager", "worker", "viewer")
def product_list_pdf(request):
    products = get_filtered_products(request).order_by("code")

    html_string = render_to_string(
        "inventory/pdf/product_list_pdf.html",
        {
            "products": products,
            "generated_at": timezone.localtime(),
        },
    )

    pdf = HTML(
        string=html_string,
        base_url=request.build_absolute_uri("/"),
    ).write_pdf()

    response = HttpResponse(
        pdf,
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        'inline; filename="filtered-products.pdf"'
    )

    return response
